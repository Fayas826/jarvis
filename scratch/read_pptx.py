import os
import sys

try:
    import pptx
except ImportError:
    os.system("pip install python-pptx")
    import pptx

def inspect_pptx(filepath):
    print(f"==================================================")
    print(f"FILE: {filepath}")
    print(f"==================================================")
    if not os.path.exists(filepath):
        print(f"NOT FOUND: {filepath}")
        return
    prs = pptx.Presentation(filepath)
    for idx, slide in enumerate(prs.slides, 1):
        print(f"\n--- Slide {idx} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    print(text)

if __name__ == "__main__":
    files = [
        r"c:\jarvis AI\jarvis\CogAgent Seminar Presentation.pptx",
        r"c:\jarvis AI\jarvis\CognitiveOS_Zeroth_FINAL_Official_Headings (1).pptx",
        r"c:\jarvis AI\jarvis\amrita_template (1).pptx"
    ]
    for f in files:
        inspect_pptx(f)
